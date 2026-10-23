#!/usr/bin/env python3
"""Конвертирует прайс prais_fotopechat.xlsx в public/prices.json.
Запуск: python3 tools/xlsx_to_json.py путь/к/prais_fotopechat.xlsx
Используется при обновлении прайса (допродажа «обновление прайса»)."""
import json, sys, os
from openpyxl import load_workbook

FMT_CODE = lambda s: s.replace('×', 'x').replace(' ', '')
PAPER = {'Глянцевая': 'glossy', 'Матовая': 'matte'}

def main(src, dst):
    wb = load_workbook(src, data_only=True)
    def rows(name):
        ws = wb[name]; it = ws.iter_rows(values_only=True); head = next(it)
        return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]

    formats = [{'code': FMT_CODE(r['Формат']), 'label': r['Формат'], 'size': r['Размер, см'],
                'express': str(r['Печать за 2 часа']).strip().lower() == 'да'} for r in rows('Форматы')]
    prices = {}
    for r in rows('Прайс'):
        p = r['Цена за 1 фото, ₽']
        if not isinstance(p, (int, float)) or p <= 0 or int(p) != p:
            raise SystemExit(f"Ошибка в прайсе: «{r['Формат']} / {r['Бумага']}» — цена должна быть целым числом > 0, сейчас: {p!r}")
        prices.setdefault(FMT_CODE(r['Формат']), {})[PAPER[r['Бумага']]] = int(p)
    for f in formats:
        if set(prices.get(f['code'], {})) != set(PAPER.values()):
            raise SystemExit(f"Ошибка в прайсе: для формата {f['label']} нужны цены на обе бумаги")
    urgency = []
    for r in rows('Срочность'):
        k = float(r['Коэффициент'])
        urgency.append({'code': r['Код'], 'label': r['Вариант'], 'term': r['Срок в калькуляторе'],
                        'percent': round(k * 100)})
    st = {r['Параметр']: r['Значение'] for r in rows('Настройки')}
    def hhmm(v):
        return v if isinstance(v, str) else v.strftime('%H:%M')
    stop = str(st['express_stop_before_close'])  # «2:30»
    sh, sm = map(int, stop.split(':')[:2])
    settings = {
        'discount_threshold': int(st['discount_threshold']),
        'discount_percent': round(float(st['discount_percent']) * 100) if float(st['discount_percent']) < 1 else int(st['discount_percent']),
        'hint_window': int(st['hint_window']),
        'min_order': int(st['min_order_rub']),
        'max_qty': int(st['max_qty_per_line']),
        'max_lines': 20,
        'express_hours': int(st['express_hours']),
        'express_stop_before_close_min': sh * 60 + sm,
        'express_start_not_before': hhmm(st['express_start_not_before']),
        'open': '10:00', 'close': '20:00', 'ready_from': '12:00',
    }
    points = [{'code': r['Код'], 'name': r['Название'], 'address': r['Адрес'], 'hours': r['Часы работы']} for r in rows('Пункты выдачи')]
    out = {'updated': __import__('datetime').date.today().isoformat(), 'formats': formats,
           'papers': [{'code': v, 'label': k} for k, v in PAPER.items()],
           'edges': [{'code': 'crop', 'label': 'С обрезкой'}, {'code': 'fit', 'label': 'С полями'}],
           'prices': prices, 'urgency': urgency, 'settings': settings, 'points': points}
    json.dump(out, open(dst, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('OK →', dst)

if __name__ == '__main__':
    here = os.path.dirname(os.path.abspath(__file__))
    main(sys.argv[1], os.path.join(here, '..', 'public', 'prices.json'))
